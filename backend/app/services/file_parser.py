"""Plain text / docx file parser for meeting and knowledge-base materials."""
from __future__ import annotations

from pathlib import Path


def parse_file(path: str | Path) -> str:
    """Return the textual content of a .txt or .docx file."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"文件不存在：{p}")
    suffix = p.suffix.lower()
    if suffix == ".txt":
        return _read_txt(p)
    if suffix == ".docx":
        return _read_docx(p)
    raise ValueError(f"不支持的文件类型：{suffix}")


def _read_txt(p: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return p.read_text(encoding=encoding).strip()
        except UnicodeDecodeError:
            continue
    return p.read_bytes().decode("utf-8", errors="ignore").strip()


def _read_docx(p: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"未安装 python-docx：{exc}") from exc
    doc = Document(str(p))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                parts.append(row_text)
    return "\n".join(parts).strip()


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 80) -> list[str]:
    """Sliding-window character chunker for embedding."""
    text = (text or "").strip()
    if not text:
        return []
    if chunk_size <= 0:
        return [text]
    if overlap < 0:
        overlap = 0
    if overlap >= chunk_size:
        overlap = chunk_size // 2

    chunks: list[str] = []
    start = 0
    n = len(text)
    step = chunk_size - overlap
    while start < n:
        end = min(start + chunk_size, n)
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        start += step
    return chunks
