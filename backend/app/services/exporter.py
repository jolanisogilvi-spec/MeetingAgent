"""Word document exporter for meeting minutes (python-docx).

Renders a meeting to a .docx with this structure:
  - Title: "MeetingAgent 会议纪要" (centered, bold, 20pt)
  - Meta block: 会议名称 / 会议时间 / 所属部门 / 参会人员 / 会议类型
  - 一、会议摘要 (heading 1)
  - 二、会议议题 (heading 1)
  - 三、关键讨论点 (heading 1, numbered list)
  - 四、决策结论 (heading 1, numbered list)
  - 五、待办事项 (heading 1, Table Grid: 责任人 | 任务内容 | 截止时间)
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import Session

from ..config import EXPORT_DIR
from ..models.department import Department
from ..models.meeting import Meeting
from ..models.person import Person

DOC_TITLE = "MeetingAgent 会议纪要"
SECTION_TITLES = (
    "一、会议摘要",
    "二、会议议题",
    "三、关键讨论点",
    "四、决策结论",
    "五、待办事项",
)
ACTION_TABLE_HEADERS = ("责任人", "任务内容", "截止时间")


def export_meeting_to_docx(
    meeting: Meeting,
    db: Session | None = None,
    output_dir: Path | None = None,
) -> Path:
    """Render the meeting to a .docx file and return its path."""
    target_dir = Path(output_dir) if output_dir else EXPORT_DIR
    target_dir.mkdir(parents=True, exist_ok=True)

    department_name = _resolve_department_name(db, meeting.department_id)
    participant_names = _resolve_participant_names(db, meeting.participant_ids or [])

    document = Document()
    _set_default_font(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(DOC_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(20)

    _add_meta_block(
        document,
        meeting=meeting,
        department_name=department_name,
        participant_names=participant_names,
    )

    data = meeting.meeting_json if isinstance(meeting.meeting_json, dict) else {}

    document.add_heading(SECTION_TITLES[0], level=1)
    _add_paragraph(document, (meeting.summary or "").strip() or "（暂无摘要）")

    document.add_heading(SECTION_TITLES[1], level=1)
    _add_paragraph(document, str(data.get("Topic") or "").strip() or "（未提取议题）")

    document.add_heading(SECTION_TITLES[2], level=1)
    _add_numbered_list(
        document,
        _as_str_list(data.get("KeyPoints")),
        empty_text="（未提取关键讨论点）",
    )

    document.add_heading(SECTION_TITLES[3], level=1)
    _add_numbered_list(
        document,
        _as_str_list(data.get("Decisions")),
        empty_text="（未提取决策结论）",
    )

    document.add_heading(SECTION_TITLES[4], level=1)
    _add_action_table(document, _normalize_action_items(data.get("ActionItems")))

    target = target_dir / _build_filename(meeting)
    document.save(str(target))
    return target


def _set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)


def _add_meta_block(
    document: Document,
    *,
    meeting: Meeting,
    department_name: str,
    participant_names: list[str],
) -> None:
    rows: list[tuple[str, str]] = [("会议名称", meeting.title or "")]
    if meeting.meeting_time:
        rows.append(("会议时间", meeting.meeting_time))
    if department_name:
        rows.append(("所属部门", department_name))
    if participant_names:
        rows.append(("参会人员", "、".join(participant_names)))
    if meeting.meeting_type:
        rows.append(("会议类型", meeting.meeting_type))

    for label, value in rows:
        para = document.add_paragraph()
        run = para.add_run(f"{label}：{value}")
        run.font.size = Pt(10)


def _add_paragraph(document: Document, text: str) -> None:
    para = document.add_paragraph()
    para.add_run(text)


def _add_numbered_list(document: Document, items: list[str], *, empty_text: str) -> None:
    if not items:
        document.add_paragraph(empty_text)
        return
    for item in items:
        document.add_paragraph(item, style="List Number")


def _add_action_table(document: Document, items: list[dict[str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(ACTION_TABLE_HEADERS))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for cell, label in zip(header_cells, ACTION_TABLE_HEADERS):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        run = para.add_run(label)
        run.bold = True

    if not items:
        row = table.add_row().cells
        row[0].text = ""
        row[1].text = "（未提取待办事项）"
        row[2].text = ""
        return

    for item in items:
        row_cells = table.add_row().cells
        row_cells[0].text = item.get("Who") or ""
        row_cells[1].text = item.get("What") or ""
        row_cells[2].text = item.get("When") or ""


def _resolve_department_name(db: Session | None, department_id: str | None) -> str:
    if db is None or not department_id:
        return ""
    dept = db.get(Department, department_id)
    return dept.name if dept else ""


def _resolve_participant_names(db: Session | None, ids: list[str]) -> list[str]:
    if db is None or not ids:
        return []
    found = db.query(Person).filter(Person.id.in_(ids)).all()
    by_id = {p.id: p.name for p in found}
    return [by_id[i] for i in ids if i in by_id]


def _as_str_list(value: Any) -> list[str]:
    if not value:
        return []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    if isinstance(value, str):
        text = value.strip()
        return [text] if text else []
    return []


def _normalize_action_items(value: Any) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    if not isinstance(value, list):
        return items
    for entry in value:
        if isinstance(entry, dict):
            items.append(
                {
                    "Who": str(entry.get("Who") or "").strip(),
                    "What": str(entry.get("What") or "").strip(),
                    "When": str(entry.get("When") or "").strip(),
                }
            )
        elif isinstance(entry, str) and entry.strip():
            items.append({"Who": "", "What": entry.strip(), "When": ""})
    return items


_FORBIDDEN_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]')


def _build_filename(meeting: Meeting) -> str:
    """Use the meeting title if available, fall back to meeting_<id>.docx."""
    raw = (meeting.title or "").strip()
    cleaned = _FORBIDDEN_FILENAME_CHARS.sub("_", raw).strip().rstrip(".")
    cleaned = cleaned[:80]
    if not cleaned:
        cleaned = f"meeting_{meeting.id}"
    return cleaned + ".docx"
